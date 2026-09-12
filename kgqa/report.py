from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .common import ROOT, read_json


def set_font(run, east_asia='宋体', latin='Times New Roman', size=10.5, bold=False):
    run.font.name = latin; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east_asia)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr(); element = OxmlElement('w:shd'); element.set(qn('w:fill'), fill); props.append(element)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers)); table.style = 'Table Grid'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]; cell.text = str(header); shade(cell, '1F4E78'); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9.2, bold=True); run.font.color.rgb = RGBColor(255, 255, 255)
    for row_number, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value); cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_number % 2: shade(cells[index], 'F3F7FA')
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 or len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            for run in cells[index].paragraphs[0].runs: set_font(run, size=8.8)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths): row.cells[index].width = Cm(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def paragraph(document, text='', bold_lead=None):
    item = document.add_paragraph(); item.paragraph_format.line_spacing = 1.35; item.paragraph_format.space_after = Pt(6)
    item.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead: set_font(item.add_run(bold_lead), bold=True)
    set_font(item.add_run(text)); return item


def add_figure(document, path, caption, number, width=15):
    if not path.exists(): return False
    document.add_picture(str(path), width=Cm(width)); document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; document.paragraphs[-1].paragraph_format.keep_with_next = True
    cap = document.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after = Pt(8)
    set_font(cap.add_run(f'图 {number} {caption}'), size=9.5)
    return True


def metric_rows(run_ids):
    rows = []
    for run_id in run_ids:
        path = ROOT / 'runs' / run_id / 'metrics_test.json'
        if not path.exists(): continue
        metric = read_json(path)
        selection = ROOT / 'runs' / run_id / 'selection.json'
        minutes = read_json(selection).get('training_seconds', 0) / 60 if selection.exists() else 0
        rows.append([run_id, f"{metric['entity_exact']:.2f}", f"{metric['entity_macro_f1']:.2f}", f"{metric['bleu']:.2f}",
                     f"{metric['rouge_l']:.2f}", f"{metric.get('latency_mean_ms', 0):.2f}", f"{metric.get('latency_p95_ms', 0):.2f}", f"{minutes:.1f}" if minutes else '-'])
    return rows


def build(output_path):
    manifest = read_json(ROOT / 'data' / 'processed' / 'manifest.json')
    environment = read_json(ROOT / 'output' / 'environment.json')
    selection = read_json(ROOT / 'output' / 'selected_model.json')
    selected = selection['selected_run']
    metrics = {run: read_json(ROOT / 'runs' / run / 'metrics_test.json') for run in ('A0', 'B0', 'B1', 'B2', 'B3', 'C0')}
    configs = {run: read_json(ROOT / 'runs' / run / 'config.json') for run in ('B0', 'B1', 'B2', 'B3', 'C0')}
    selections = {run: read_json(ROOT / 'runs' / run / 'selection.json') for run in ('B0', 'B1', 'B2', 'B3', 'C0')}

    document = Document(); section = document.sections[0]
    section.top_margin = Cm(2.1); section.bottom_margin = Cm(2.1); section.left_margin = Cm(2.3); section.right_margin = Cm(2.3)
    styles = document.styles
    for name, size in [('Normal', 10.5), ('Title', 20), ('Heading 1', 15), ('Heading 2', 12)]:
        style = styles[name]; style.font.name = 'Times New Roman'; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体' if name != 'Normal' else '宋体')
        if name == 'Title':
            ppr = style._element.get_or_add_pPr()
            for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
    styles['Heading 1'].paragraph_format.space_before = Pt(14); styles['Heading 1'].paragraph_format.space_after = Pt(8)
    styles['Heading 2'].paragraph_format.space_before = Pt(10); styles['Heading 2'].paragraph_format.space_after = Pt(6)

    title = document.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for border in list(title._p.get_or_add_pPr().findall(qn('w:pBdr'))): title._p.get_or_add_pPr().remove(border)
    set_font(title.add_run('基于 Wikidata 电影知识图谱的问答系统与答案生成实验报告'), east_asia='黑体', size=20, bold=True)
    intro = document.add_paragraph(); intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(intro.add_run('人工智能 I 课程实验'), east_asia='黑体', size=12)
    paragraph(document, '本报告记录电影单跳知识图谱问答实验的实际数据构建、模型实现、参数调整、统一评测和结果分析。实验比较可解释的管道式方法与融合 TransE 结构向量的 FLAN-T5 生成方法，并通过学习率对照和嵌入消融判断主要性能变化。')

    document.add_heading('二 实验基本信息', level=1)
    add_table(document, ['项目', '内容'], [['实验名称', '基于知识图谱的问答系统与答案生成'], ['姓名', '待填写'], ['学号', '待填写'], ['班级', '待填写'], ['实验完成日期', '2026年9月12日']], [4, 11])

    document.add_heading('三 实验环境', level=1)
    add_table(document, ['项目', '实际配置'], [['操作系统', environment['operating_system']], ['Python 路径', environment['python_executable']],
              ['Python 版本', environment['python_version'].split('|')[0].strip()], ['CPU', f"{environment.get('processor') or '12th Gen Intel Core i7-12700H'}，{environment['logical_cpu_count']} 个逻辑处理器"],
              ['内存', f"{environment['memory_gb']} GB"], ['GPU', environment['gpu']], ['CUDA', f"可用，PyTorch CUDA {environment['torch_cuda']}"],
              ['主要依赖', ', '.join(f'{k} {v}' for k, v in environment['packages'].items())]], [4, 11])
    add_figure(document, ROOT / 'output' / 'evidence' / 'environment.png', '终端环境配置记录', 1, 15.5)

    document.add_heading('四 数据集与实验设置', level=1)
    paragraph(document, f"数据来自 Wikidata 的电影子图，使用导演 P57、类型 P136、制作国家 P495 和原始语言 P364 四类关系。冻结快照的 SHA-256 为 {manifest['snapshot_sha256']}，包含 {manifest['movies']} 部电影、{manifest['entities']} 个实体、{manifest['triples']} 条三元组和 {manifest['qa']} 条问答。")
    split = manifest['qa_by_split']
    add_table(document, ['数据项', '数量或设置'], [['训练集问答', split['train']], ['验证集问答', split['valid']], ['测试集问答', split['test']],
              ['电影实体划分', '700 / 150 / 150'], ['随机种子', 42], ['划分原则', '按电影 QID 分组，同一电影的原问题和改写位于同一集合']], [5, 10])
    paragraph(document, 'A0 采用电影标题与别名词典完成实体链接，使用词和字符 TF-IDF 特征及逻辑回归判断关系，再查询冻结子图并通过模板生成完整句。B0 至 B3 采用 FLAN-T5-small，以问题和相关三元组文本为输入，并将冻结的 TransE 头实体、关系和尾实体向量投影为结构前缀。C0 保留相同三元组文本，但移除结构向量。')

    document.add_heading('五 实验结果', level=1)
    document.add_heading('5.1 基础实验设置与运行记录', level=2)
    paragraph(document, '正式实验固定使用上述数据划分。A0 的关系置信度阈值在验证集上选择；所有神经模型均从同一 FLAN-T5-small 初始权重独立训练，使用随机种子 42，测试集只在验证选模完成后评估一次。')
    b0 = configs['B0']
    add_table(document, ['模块', '主要参数'], [
        ['A0', f"TF-IDF + 逻辑回归；验证选择阈值 {read_json(ROOT / 'runs' / 'A0' / 'config.json')['threshold']:.2f}"],
        ['TransE', '维度 100；学习率 0.001；batch size 256；100 epoch；每个正例 5 个负样本'],
        ['B0', f"学习率 {b0['learning_rate']:.0e}；{b0['epochs']} epoch；batch size {b0['batch_size']}；有效 batch size {b0['effective_batch']}；BF16 自动混合精度"],
        ['选模指标', '验证集答案实体 Macro-F1']],[4, 11])
    figure_number = 2
    for filename, caption in [('training_b0.png', 'B0 模型训练过程记录'), ('evaluation.png', '测试评估与参数实验运行记录')]:
        if add_figure(document, ROOT / 'output' / 'evidence' / filename, caption, figure_number, 15.5): figure_number += 1

    document.add_heading('5.2 基础实验结果', level=2)
    add_table(document, ['实验', '实体准确率', '实体 F1', 'BLEU', 'ROUGE-L', '平均延迟 ms', 'P95 延迟 ms', '训练 min'], metric_rows(['A0', 'B0']), [1.4, 2, 1.8, 1.6, 1.8, 2.2, 2.2, 1.8])
    a0, b0m = metrics['A0'], metrics['B0']
    paragraph(document, f"A0 的实体集合准确率为 {a0['entity_exact']:.2f}%，Macro-F1 为 {a0['entity_macro_f1']:.2f}%，平均延迟为 {a0['latency_mean_ms']:.2f} ms。B0 的实体集合准确率为 {b0m['entity_exact']:.2f}%，Macro-F1 为 {b0m['entity_macro_f1']:.2f}%，BLEU 为 {b0m['bleu']:.2f}，ROUGE-L 为 {b0m['rouge_l']:.2f}。B0 文本重叠指标较高，但多答案问题存在漏答，使集合完全匹配指标低于实体 F1。")

    document.add_heading('5.3 参数修改与重复实验', level=2)
    rows = []
    for run in ('B0', 'B1', 'B2', 'B3', 'C0'):
        cfg, sel = configs[run], selections[run]
        linked_valid_path = ROOT / 'runs' / run / 'metrics_valid.json'
        valid_f1 = read_json(linked_valid_path)['entity_macro_f1'] if linked_valid_path.exists() else sel['best_valid_f1']
        rows.append([run, f"{cfg['learning_rate']:.0e}", '使用' if cfg['use_embeddings'] else '移除', sel['best_epoch'], f"{valid_f1:.2f}", f"{sel['training_seconds']/60:.1f}"])
    add_table(document, ['实验', '学习率', 'TransE', '最佳 epoch', '验证 F1', '训练 min'], rows, [2, 2.5, 2.5, 2.5, 2.5, 3])
    paragraph(document, 'B0 至 B3 仅改变学习率，数据、训练轮数、随机种子、模型初始权重和其他超参数保持一致。C0 与 B0 使用相同学习率和三元组文本，仅移除 TransE 结构前缀。')

    document.add_heading('5.4 基础实验与参数修改实验的结果对比', level=2)
    add_table(document, ['实验', '实体准确率', '实体 F1', 'BLEU', 'ROUGE-L', '平均延迟 ms', 'P95 延迟 ms', '训练 min'], metric_rows(['B0', 'B1', 'B2', 'B3', 'C0']), [1.4, 2, 1.8, 1.6, 1.8, 2.2, 2.2, 1.8])
    figures = [('learning_rate_loss.png', '不同学习率下的验证损失曲线'), ('learning_rate_metrics.png', '学习率与最终测试指标'),
               ('embedding_ablation.png', 'TransE 结构向量消融对比'), ('robustness.png', 'A0 与最佳神经系统的鲁棒性对比')]
    for filename, caption in figures:
        if add_figure(document, ROOT / 'output' / 'figures' / filename, caption, figure_number, 11.5): figure_number += 1

    document.add_heading('5.5 结果对比分析', level=2)
    lr_metrics = {run: metrics[run]['entity_macro_f1'] for run in ('B0', 'B1', 'B2', 'B3')}
    best_test = max(lr_metrics, key=lr_metrics.get); worst_test = min(lr_metrics, key=lr_metrics.get)
    paragraph(document, f"参数修改后，四组学习率的测试实体 Macro-F1 分布在 {min(lr_metrics.values()):.2f}% 至 {max(lr_metrics.values()):.2f}% 之间。按验证集 F1 选择的模型为 {selected}，其验证 F1 为 {selection['validation_f1']:.2f}%；若只观察测试集，最高的是 {best_test}（{lr_metrics[best_test]:.2f}%），最低的是 {worst_test}（{lr_metrics[worst_test]:.2f}%）。模型选择严格依据验证集，未使用测试集反向挑选参数。")
    c0 = metrics['C0']
    delta = b0m['entity_macro_f1'] - c0['entity_macro_f1']
    direction = '提高' if delta >= 0 else '降低'
    paragraph(document, f"嵌入消融中，B0 的实体 Macro-F1 为 {b0m['entity_macro_f1']:.2f}%，C0 为 {c0['entity_macro_f1']:.2f}%。加入 TransE 后该指标{direction} {abs(delta):.2f} 个百分点。由于两组仍保留相同三元组文本，这一差值反映结构向量在本次单种子实验中的附加影响。")
    robust_a0 = read_json(ROOT / 'runs' / 'A0' / 'metrics_robustness.json')
    robust_n = read_json(ROOT / 'runs' / selected / 'metrics_robustness.json')
    shared = [v for v in ('clean', 'paraphrase', 'alias', 'typo', 'unanswerable') if v in robust_a0 and v in robust_n]
    robust_text = '；'.join(f"{v}：A0 {robust_a0[v]['entity_macro_f1']:.2f}%，{selected} {robust_n[v]['entity_macro_f1']:.2f}%" for v in shared)
    paragraph(document, f"鲁棒性结果为：{robust_text}。实体链接错误、关系改写覆盖不足和拼写噪声是主要下降来源；生成系统还可能出现多答案漏答。学习率过低会限制五轮内的参数更新，过高则可能使收敛波动增大，因此最终结果需要结合验证损失曲线和实体 F1 共同判断。")
    error_path = ROOT / 'output' / 'error_analysis.json'
    if error_path.exists():
        errors = read_json(error_path)
        labels = {'entity_linking_error': '实体链接错误', 'relation_error': '关系识别错误', 'missed_answer': '漏答', 'incomplete_answer': '多答案不完整', 'unsupported_answer_entity': '无依据答案实体', 'failed_refusal': '拒答失败'}
        parts = []
        for run in ('A0', selected):
            counts = errors[run]['counts']
            detail = '，'.join(f"{labels.get(k, k)} {v} 题" for k, v in counts.items() if k != 'correct' and v)
            parts.append(f"{run}：{detail or '未检出错误'}")
        paragraph(document, '逐题错误归因结果为：' + '；'.join(parts) + '。该归因按实体链接、关系识别、漏答、答案不完整、无依据实体和拒答失败的顺序判定。')
    paragraph(document, '本次结果总体符合预期：A0 在固定单跳事实查询上具有较高精确性和极低延迟，神经模型的语言生成质量较好，但严格集合匹配会受到漏答影响。结论只适用于本次固定知识库、单次随机种子和电影单跳任务，不推断模型能够回答知识库之外的未知事实。')
    paragraph(document, '人工评价表已生成，准确性、完整性、流畅度和无依据生成字段必须由实验者逐题核查后填写；本报告不以自动指标代替人工事实评价。')

    document.add_heading('六 实验总结', level=1)
    paragraph(document, f"本实验完成了 Wikidata 数据冻结、按电影实体划分、A0 管道问答、TransE 嵌入、五组 FLAN-T5 训练、统一评测、鲁棒性测试和嵌入消融。实际结果表明，{selected} 在验证集上得到最高实体 F1；A0 在固定事实查询上更容易保持集合精确匹配和低延迟。实验也说明，文本相似度高并不保证多答案实体集合完全正确，知识图谱问答应同时报告实体级指标、文本指标、拒答能力、鲁棒性和运行成本。")
    paragraph(document, '参考资料包括 Wikidata 数据访问文档、TransE 原始论文、PyKEEN 文档、FLAN-T5 模型卡、WebNLG、SacreBLEU 和 ROUGE 论文。完整数据查询和实验设计见项目文件《实验设计与参考资料》。')
    document.add_heading('七 参考资料', level=1)
    references = [
        '[1] Wikidata. Wikidata Query Service User Manual. https://www.wikidata.org/wiki/Wikidata:SPARQL_query_service/Wikidata_Query_Help',
        '[2] Wikidata. Data access and entity data documentation. https://www.wikidata.org/wiki/Wikidata:Data_access',
        '[3] Bordes A, Usunier N, Garcia-Duran A, et al. Translating Embeddings for Modeling Multi-relational Data. NeurIPS, 2013.',
        '[4] Ali M, Berrendorf M, Hoyt C T, et al. PyKEEN 1.0: A Python Library for Training and Evaluating Knowledge Graph Embeddings. JMLR, 2021.',
        '[5] Google. FLAN-T5-small model card. https://huggingface.co/google/flan-t5-small',
        '[6] Gardent C, Shimorina A, Narayan S, Perez-Beltrachini L. Creating Training Corpora for NLG Micro-Planners. ACL, 2017.',
        '[7] Post M. A Call for Clarity in Reporting BLEU Scores. WMT, 2018.',
        '[8] Lin C Y. ROUGE: A Package for Automatic Evaluation of Summaries. ACL Workshop, 2004.'
    ]
    for item in references:
        p = document.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.5); p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(item), size=9.5)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True); document.save(output_path)


if __name__ == '__main__': build(ROOT / 'output' / '人工智能I_实验报告.docx')

