"""对统一模板报告作PR12审计勘误，保留原图与原生公式，重复执行不重复插入。"""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
FILES = ['个人实验报告.docx', '夏俊杰_实验报告.docx', '双人分工与复现实验计划.docx']
ENV_NOTE = ('本节区分原主实验、H盘本机复核和V100补充运行。2026年10月31日已保存本机及服务器真实终端文本，'
            '记录见verification/2026-10-31/pr12_audit_fixes/。当前交付尚无符合要求的真实终端环境截图，'
            '浏览器运行图不能替代该项课程材料；两台机器不要求使用相同版本。')
MEMORY_NOTE = ('GPU1上本账户P0服务PID的三次nvidia-smi采样均为7542 MiB（约7.37 GiB）；'
               '这是2026年10月31日的稳定进程采样，不是全生命周期峰值。原2.1—2.5 GiB说法撤回；'
               '25%仅为PyTorch分配器上限，不含全部CUDA开销，也不代表整卡独占。')
FREEFORM_NOTE = ('四关系自由问法v2微调已完成3轮（学习率3e-5、有效batch16、seed42），按验证F1选择第3轮。'
                 '验证总体F1由89.65%到93.20%，可答F1由88.03%到99.53%，拒答却从95.47%降至70.53%。'
                 '6904题混合测试实体F1为69.4081%、可答F1为72.48%、拒答58.33%。其中738题的国家模板已出现于原训练题，'
                 '不能把全部测试称为未见句式。没有旧B0同测试集对照，不能与原模板测试直接作提升差值。'
                 '总耗时15632.05秒，进程采样峰值5600 MiB；未替换界面默认模型。')
HUMAN_NOTE = ('课程任务书第4项明确要求人工评价准确性、完整性和流畅度。当前匿名样本已准备，尚无完整真人评分结果，'
              '因此本报告不声称该项要求已完成。至少一名真实评价者对相同100题的A0/B0共200份答案评分；'
              '准确性与完整性各0—2、流畅度1—5、无依据断言0/1。记录实际人数；多人独立评分时另报一致性，'
              '单人不计算一致性。空白不当零分，不用自动指标或LLM代填。')
ASSET_NOTE = ('冻结V2图谱及自由问法v2数据、逐题预测、验证选中检查点和配套TransE已另行发布，'
              '清单为configs/submission-assets-20261031.json，下载页为GitHub Release submission-assets-20261031。'
              '按清单恢复并逐文件校验；重采Wikidata属于新版本。检查点支持推理，不含优化器状态，不能声称可精确续训。'
              '课程代码包继续排除完整数据、大型权重与缓存。')
GUARD_NOTE = ('2026年10月31日审计修订后，P0最终展示始终由render_tool按执行结果生成完整受控证据句，模型原文仅存于raw_answer。'
              '只有与受控句仅空白不同且引用完全匹配的文本可标为accepted；正确引用不能证明任意附加断言成立。'
              '中英文合成控制器测试覆盖导演误写为演员、主客体颠倒及追加无证据获奖信息，均不得进入最终答案。'
              '这属于控制器测试，不是声称真实Qwen生成了这些错误。错误计划及图谱缺失仍需分别检查。')
REPLAY_NOTE = ('对原20题保存输出的离线控制器重放使模板回退从11/20变为13/20，原accepted中的2条文本被拦下；'
               '没有调用模型、重新执行计划或重新测量延迟。上表11/20及延迟保留原开发运行口径，'
               '新重放另存pr12_audit_fixes/p0_guard_replay/，不能解释为模型能力提升。')


def paragraph_before(doc, anchor, text, style=None):
    p = doc.add_paragraph(text, style)
    anchor._p.addprevious(p._p)
    return p


def correct_report(doc, *, plan=False):
    # 原图、公式所在段落不赋值，避免破坏其OOXML。
    for p in doc.paragraphs:
        text = p.text
        if text.startswith('环境截图待补'):
            p.text = ENV_NOTE
        elif 'GPU1实测进程约2.1—2.5GiB' in text:
            p.text = text.replace('GPU1实测进程约2.1—2.5GiB；25%是PyTorch分配器上限，不代表整卡独占。', MEMORY_NOTE)
        elif text.startswith('四关系自由问法v2微调已完成'):
            p.text = FREEFORM_NOTE
        elif text.startswith('人工评分仍为空'):
            p.text = ('本报告完成日期为2026年10月30日；10月31日扩展与审计按真实日期单列。人工评分和终端环境截图尚未完成，'
                      '相关课程要求不能记为已满足。评分与环境采集入口见本轮修订说明。')
        elif text.startswith('本机完整代码回归716项通过'):
            p.text = ('本轮审计修订的完整代码回归740项通过（37.94秒），Ruff检查通过。协议模拟、Schema、执行器及Web测试验证工程行为，'
                      '不能代替真实模型自然语言能力评价。旧716项运行记录保留原日期及上下文。')
        elif text == '待补材料':
            p.text = '评价协议与实验局限'
        elif text.startswith(('两张实际浏览器模型测试截图已移入', '图5-1、图5-2记录实际浏览器模型测试。')):
            p.text = '图5-1、图5-2记录实际浏览器模型测试；人工评价协议、资产获取和当前材料局限见前文，未以自动指标替代真人评分。'
        elif text.startswith('提交前仍需补入真实终端环境截图') or (plan and text.startswith('本节区分原主实验')):
            p.text = '人工评分与终端环境截图仍是未完成的课程材料。两人核对原始记录后更新报告，不把评分入口或终端文本记作评分结果或截图。'
        elif text.startswith('人工计划采用相同100题'):
            p.text = ('评分入口为python -m kgqa.human_review_web --port 7865，浏览器访问http://127.0.0.1:7865/。'
                      '填写评价者代号后可续填、保存及导出；原始记录保存在output/human_review/。完成后使用kgqa.human_eval汇总，'
                      '保留原始评分和评价者人数，不把空白表当已完成的人工实验。')
        elif text.startswith('6 待完成事项与改进顺序'):
            p.text = '6 课程材料状态与改进顺序'
        elif '人工评分是本项目原定评价计划' in text:
            p.text = HUMAN_NOTE
        elif '本轮结果以保持固定图谱' in text:
            p.text = text.replace('真实人工评分仍待完成', '人工评价尚无完整真人评分，课程该项要求尚未完成')
    for table in doc.tables:
        for row in table.rows:
            for c in row.cells:
                if c.text == '自设计划尚未评分，不虚构分数':
                    c.text = '任务书必需项；评分入口已准备，尚无完整真人评分'
                elif c.text == '本机716 passed，39.21秒；协议模拟、执行器、Web和回归分别验证':
                    c.text = '审计修订740 passed，37.94秒；工程测试不等于模型成绩'
                elif '3轮已完成；验证F1 93.20%' in c.text:
                    c.text = '3轮完成；混合6904题F1 69.4081%；738题已见模板100%，其余6166题65.7466%'
    # 夏报告末尾已有完整的评价局限小节，删除旧稿遗留的重复编辑段。
    for p in list(doc.paragraphs):
        if p.text == '评价协议与实验局限' or p.text.startswith('图5-1、图5-2记录实际浏览器模型测试；人工评价协议'):
            p._element.getparent().remove(p._element)
    if any(p.text == '自由问法模板审计勘误' for p in doc.paragraphs):
        return
    if plan:
        end = next(p for p in doc.paragraphs if p.text.startswith('4 复现顺序'))
    else:
        end = next(p for p in doc.paragraphs if p.text.startswith('5.5 结果对比分析'))
    paragraph_before(doc, end, '自由问法模板审计勘误', 'Heading 3')
    paragraph_before(doc, end, '审计比较实际原训练题与新增训练题的合并集合，以片名占位、NFKC、大小写及标点归一后的精确模板匹配分组。'
                     '电影划分未发现交叉；family名不同并不保证句式未见。以下由原保存预测真实重计分，无重新训练、生成或选模。')
    from unify_course_reports import table_before
    table_before(doc, end, ['测试分组', '题数', '实体Macro F1'], [
        ['已见规范化模板', '738', '100.0000%'],
        ['其余模板', '6166', '65.7466%'],
        ['混合全体', '6904', '69.4081%'],
    ], [82, 32, 52])
    paragraph_before(doc, end, '738+6166=6904。其余模板仅表示未找到上述规范化精确匹配，不等于独立真人自然问答；'
                     '6868道验证题的该类重叠为0。原v1/v2冻结数据和原分数保留，v3新增合并训练模板防重叠校验，未据此重新训练。'
                     '原预测、分组题号和复算脚本见verification/2026-10-31/pr12_audit_fixes/freeform/及scripts/audit_freeform_overlap.py。')
    paragraph_before(doc, end, 'P0证据回答修订', 'Heading 3')
    paragraph_before(doc, end, GUARD_NOTE)
    paragraph_before(doc, end, REPLAY_NOTE)
    end = (next(p for p in doc.paragraphs if p.text.startswith('6 课程材料')) if plan
           else next(p for p in doc.paragraphs if p.text.startswith('6 实验总结')))
    paragraph_before(doc, end, '开发复杂度与部署成本比较', 'Heading 3')
    paragraph_before(doc, end, 'A0由名称词典、TF-IDF关系分类器、事实查询及答案模板组成，依赖与部署较轻；'
                     '增加属性和新问法时需要维护分类标签、别名及模板。B0除数据清洗外，还需要TransE映射、预训练权重、'
                     '结构投影、训练与验证选模、解码和检查点管理，显存、依赖兼容与复现环节更多。'
                     '本轮B0/C0同分说明额外结构模块并不保证收益，需结合实际结果决定维护价值。')
    paragraph_before(doc, end, 'P0不做梯度训练，但增加模型服务、候选绑定、计划Schema、确定性执行、超时处理和证据审计；'
                     '免训练减少训练环节，仍有较多开发与部署工作，原20题平均约5.95秒也增加交互成本。'
                     '比较依据为实际kgqa/pipeline.py、neural.py、semantic_qa.py等模块及运行记录，未编造个人工时或按代码行数估算开发时间。')
    paragraph_before(doc, end, '人工评价与复现材料状态', 'Heading 3')
    paragraph_before(doc, end, HUMAN_NOTE)
    paragraph_before(doc, end, ASSET_NOTE)
    if plan:
        paragraph_before(doc, end, MEMORY_NOTE)


def media(path):
    with ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--reports-dir', type=Path, default=ROOT / 'reports')
    args = parser.parse_args()
    backup = ROOT / '.work/pr12-report-backups' / datetime.now().strftime('%Y%m%d_%H%M%S')
    backup.mkdir(parents=True, exist_ok=False)
    from unify_course_reports import format_doc
    records = []
    for name, author in zip(FILES, ['周新永', '夏俊杰', '周新永 夏俊杰']):
        path = args.reports_dir / name
        shutil.copy2(path, backup / name)
        old_media = media(path)
        doc = Document(path)
        old_math = len(doc.element.xpath('.//m:oMath'))
        correct_report(doc, plan=name == FILES[2])
        format_doc(doc, author)
        if name == FILES[0]:
            next(p for p in doc.paragraphs if p.text.startswith('附录 实验记录索引')).paragraph_format.page_break_before = True
        tmp = path.with_suffix('.correcting.docx')
        doc.save(tmp)
        if media(tmp) != old_media or len(Document(tmp).element.xpath('.//m:oMath')) != old_math:
            raise ValueError('原图片或公式发生变化')
        tmp.replace(path)
        records.append({'file': path.relative_to(ROOT).as_posix(), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                        'media_preserved': True, 'native_math_count': old_math})
    output = ROOT / 'verification/2026-10-31/pr12_audit_fixes/report_edit.json'
    output.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(records, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
