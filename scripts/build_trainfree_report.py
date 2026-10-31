"""从实跑结果创建独立中文Word补充稿，不覆盖原始课程报告。"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / 'verification/2026-10-31/trainfree'
OUT = ROOT / 'reports/无需训练增强实验补充报告_20261031.docx'


def main():
    summary = json.loads((EVIDENCE / 'runtime_5000_final_v2/summary.json').read_text(encoding='utf-8'))
    manifest = json.loads((EVIDENCE / 'runtime_5000_final_v2/manifest.json').read_text(encoding='utf-8'))
    assert manifest['status'] == 'complete' and manifest['source_unchanged'] and manifest['inputs_unchanged']
    doc = Document()
    # Some bundled templates carry an inherited decorative Title border.
    for root in (doc.styles._element, doc._element):
        for border in root.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('电影知识图谱问答  补充实验  ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    def p(text):
        return doc.add_paragraph(text)

    def heading(text):
        doc.add_heading(text, 1)

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for cell, text, width in zip(t.rows[0].cells, headers, widths):
            cell.text = text; cell.width = Inches(width)
        repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell, text, width in zip(t.add_row().cells, row, widths):
                cell.text = text; cell.width = Inches(width)
        for i, row in enumerate(t.rows):
            for j, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ('top', 'left', 'bottom', 'right'):
                    b = OxmlElement(f'w:{edge}'); b.set(qn('w:val'), 'single')
                    b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
                tc.append(borders)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E8EEF3' if i == 0 else ('F7F9FB' if i % 2 == 0 else 'FFFFFF'))
                tc.append(shading)
                margins = OxmlElement('w:tcMar')
                for edge in ('top', 'left', 'bottom', 'right'):
                    m = OxmlElement(f'w:{edge}'); m.set(qn('w:w'), '90'); m.set(qn('w:type'), 'dxa'); margins.append(m)
                tc.append(margins)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(2)
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = Pt(10.5); run.bold = i == 0
        p('')

    doc.add_paragraph('电影知识图谱问答无需训练增强实验', 'Title')
    p('课程报告补充稿　2026年10月31日')
    p('本补充记录在冻结A0与B0模型外增加的双语解析、证据核验和工具回退。5000电影运行包的六组对照已完成，没有新增训练或修改模型权重。结果支持开发验收范围内的可靠性改进，不能解释为纯神经模型自由问答能力全面提升。本稿作为既有课程主实验报告的补充。')
    heading('5.1 实验目的与设计')
    p('已有系统能回答导演、电影类型、制作国家和原始语言四类事实，但自然表达可能包含片名歧义、未支持属性或额外条件，神经模型也可能漏答或生成无依据内容。本实验研究无需训练的系统增强，并比较保守拒答与工具补答的作用。')
    p('处理流程为：中英文问题 → 电影与属性解析 → 范围及歧义检查 → 标准问题 → 冻结模型 → 关系与实体核验 → 接受、拒答或显式工具回退。演员、编剧等十二属性工具属于独立扩展任务，不合并到本实验分数。')
    heading('5.2 实现方法')
    for text in (
        '复用中英文名称、别名、年份及中文目录；相似片名要求确认。否定、筛选、比较、数量和未解释条件不被静默删除。原始语言与对白语言、导演与掌镜分别判断。',
        '把明确的四关系问句转换为模型熟悉的英文表达；神经输入按属性筛选原先最多24条可见事实，文本与TransE向量同步选择，保留原有检索上限。',
        '检查所问关系的实体集合和多答案完整性，仅接受与事实一致的两种已知完整句式。其他自由表述也可能被保守拒绝，这一核验器并非通用语义事实验证模型。',
        '电影、属性与证据明确时，混合系统允许完整属性查询及模板回退。记录原始输出、模型实际证据、核验理由和回退标志；界面明确显示工具介入。'
    ):
        p(text)

    doc.add_page_break()
    heading('5.3 实验设置与结果')
    p('使用固定5000电影运行包中的A0与验证集选中的B0。对照为raw原始入口、guarded增强核验但不补答、hybrid增强并允许工具回退。raw保留仓库原有中文适配。运行前后绑定的数据、权重及源码哈希一致。实际环境为Windows、Python3.11.7、PyTorch2.6.0+cu124、RTX3070Ti Laptop GPU，CPU线程2。')
    p('48道程序编写的开发验收题含24道可答、12道不可答、12道需澄清题，中英问句各24道。可答题只涉及6部电影，参考实体由运行前声明的电影与关系从固定KB读取；16题电影来自原train划分、8题来自test划分。中文题主要使用英文片名或QID与中文问句混合。这不是独立测试或人工评分。')
    keys = ['a0_raw', 'a0_guarded', 'a0_hybrid', 'neural_raw', 'neural_guarded', 'neural_hybrid']
    labels = ['A0原始', 'A0核验', 'A0混合', 'B0原始', 'B0核验', 'B0混合']
    metrics = ['supported_entity_f1', 'unanswerable_refusal_accuracy', 'clarification_status_accuracy', 'tool_fallback_rate']
    table(['配置', '可答实体F1', '不可答拒答', '允许澄清状态', '工具回退'],
          [[label, *[f'{summary[key][m]:.2f}%' for m in metrics]] for label, key in zip(labels, keys)],
          [1.1, 1.4, 1.4, 1.55, 1.45])
    p('可答、不可答、需澄清的分母分别是24、12、12题，工具回退分母是全部48题。B0混合系统有2题回退，占可答题8.33%和全部题4.17%。允许澄清状态匹配包含部分题允许的保守拒答，不等于实际发起澄清的比例。')
    p('A0原始、核验、混合的平均延迟分别为6.30、19.41、19.90 ms，P95为3.07、38.29、29.83 ms。B0三组平均延迟分别为886.11、711.56、691.31 ms，P95为1742.25、1816.33、1811.20 ms。完整延迟表与逐题记录保留在配套Markdown和结果JSON。')
    p('延迟为48题逐题墙钟时间，初始化另计，首次查询开销保留。A0少数首次调用会使均值大于P95。增强入口会提前拦截问题，因此全题平均时间下降不代表生成器本身加速；小样本和并行界面验证影响本轮延迟，未作稳定性能结论。')

    doc.add_page_break()
    heading('5.4 结果分析')
    p('A0增强在本组保留100%的可答实体F1，且不可答拒答率从50%升至100%，主要收益来自入口范围保护。B0严格核验的可答F1由94.79%降到91.67%，说明拦截错误生成存在覆盖代价；B0混合系统通过2次明确标记的工具回退达到本组100%，不能称为纯神经能力提高。')
    p('非可答题的自动非拒答作答代理率，A0由50%降至0%，B0由45.83%降至0%。该指标依据显式拒答字段或固定拒答句判定，其他拒答措辞需要人工复核，不能直接视为真人核查的事实错误率。')
    figure = EVIDENCE / 'figures/trainfree_quality.png'
    if figure.exists():
        doc.add_picture(str(figure), width=Inches(6.8))
        p('图1　固定开发验收中的实体F1与不可答拒答率。纵轴均为0至100，不代表全库自由问答准确率。')
    p('本轮只进行组合策略对照，未分别消融问法规范化、关系筛选和句式核验，不能判断各组件的独立贡献；也没有多seed显著性分析或新增真人评价。同资产A0复跑得到相同主要结果，只是配对核对，不是独立数据或统计交叉验证。')
    heading('5.5 总结与后续改进')
    p('无需追加训练即可增加可解释的范围保护、歧义澄清和事实保障。现有局限包括有限语法降低覆盖、中文实体名缺失、严格表面核验带来的回退，以及知识库自身可能不完整。后续应先冻结独立人工中英挑战集，再评价冻结大模型规划器、语义检索和更灵活的证据核验。Qwen规划器本轮未部署。')
    p('此前自由问法微调已经完成，其测试实体F1为69.408073%、拒答58.333333%，尚未替换当前界面模型；不能写成尚未训练，也不能与本开发验收直接相减。详细训练记录与原始主实验继续保留。')
    p('工程验证：轻量回归567项通过、2项跳过；完整神经环境针对性回归106项通过；致命代码检查通过。对应日志保存在验证目录。原报告、历史冻结数据、旧模型及其他用户任务保留。')

    doc.add_page_break()
    heading('运行证据与复现')
    p('本机启动：在H盘项目根目录执行 scripts/run_trainfree_5000.cmd，浏览器访问 http://127.0.0.1:7863/。页面默认开启自由问答增强，取消勾选可查看原入口；展开卡片可核对原始输出与工具介入。')
    for name, caption in [('中文双系统.png', '图2　5000电影双系统中文问答与事实证据。'),
                          ('工具回退与原始生成.png', '图3　多答案遗漏触发显式工具回退，原始生成可展开追溯。')]:
        image = EVIDENCE / 'screenshots' / name
        if image.exists():
            doc.add_picture(str(image), width=Inches(5.8))
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            p(caption)
    p('实跑指标、参考事实及预测：verification/2026-10-31/trainfree/runtime_5000_final_v2。完整实现、复现命令和限制说明：docs/无需训练增强与报告补充_20261031.md。')
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
