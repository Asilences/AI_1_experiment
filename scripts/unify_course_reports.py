"""按课程要求统一三份报告，保留已有图像、公式和实验表格。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FILES = ["个人实验报告.docx", "夏俊杰_实验报告.docx", "双人分工与复现实验计划.docx"]
ZHOU_SCOPE = "A0管道、电影实体链接与候选确认、关系阈值诊断、P0语义规划与确定性执行、严格评价、拒答分析及证据交互"
XIA_SCOPE = "5000电影数据与来源核验、TransE、B0—B3/C0训练、自由问法微调复核、冻结模型服务、证据生成与运行环境"
P0_LIMIT = "20题为程序编写的开发验收，中英文各10题，仅涉及Inception和Interstellar两部电影；已用于开发诊断，不是独立泛化测试。结构化判分不代表生成句的事实正确性或流畅度，尚无人工文本评分。"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def para(doc, text, style=None):
    return doc.add_paragraph(text, style=style)


def anchor(doc, starts):
    matches = [p for p in doc.paragraphs if p.text.startswith(starts)]
    if len(matches) > 1:
        headings = [p for p in matches if p.style.name.startswith('Heading')]
        if len(headings) == 1:
            return headings[0]
    if len(matches) != 1:
        raise ValueError(f"段落定位不唯一 {starts!r}: {len(matches)}")
    return matches[0]


def before(doc, target, text, style=None):
    p = doc.add_paragraph(text, style=style)
    target._p.addprevious(p._p)
    return p


def table_before(doc, target, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for c, h in zip(t.rows[0].cells, headers):
        c.text = h
    for values in rows:
        for c, v in zip(t.add_row().cells, values):
            c.text = str(v)
    target._p.addprevious(t._tbl)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Mm(w)
        for c, w in zip(t.columns, widths):
            c.width = Mm(w)
    return t


def replace(doc, old, new):
    p = anchor(doc, old)
    p.text = new
    return p


def remove(p):
    p._element.getparent().remove(p._element)


def font(r, size=10.5, bold=None):
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        r.bold = bold
    rf = r._element.get_or_add_rPr().get_or_add_rFonts()
    for a, value in [("ascii", "Times New Roman"), ("hAnsi", "Times New Roman"), ("eastAsia", "宋体")]:
        rf.set(qn(f"w:{a}"), value)


def format_doc(doc, author):
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin, s.bottom_margin = Mm(21), Mm(20)
        s.left_margin, s.right_margin = Mm(22), Mm(22)
        s.different_first_page_header_footer = False
    for name, size in [("Normal", 10.5), ("Title", 18), ("Subtitle", 10.5), ("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 11)]:
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), "宋体")
        st.paragraph_format.line_spacing = 1.25
    for p in doc.paragraphs:
        if not p.text and not p._p.xpath('.//w:drawing | .//m:oMath | .//w:bookmarkStart | .//w:sectPr'):
            remove(p)
            continue
        pf = p.paragraph_format
        pf.line_spacing = 1.25
        pf.space_before, pf.space_after = Pt(0), Pt(4)
        pf.page_break_before = False
        pf.widow_control = True
        pf.keep_together = False
        pf.keep_with_next = False
        name = p.style.name
        if name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_after, pf.keep_with_next = Pt(10), True
            size, bold = 18, True
        elif name == "Subtitle":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            size, bold = 10.5, False
        elif name.startswith("Heading"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.space_before, pf.space_after = Pt(10), Pt(5)
            pf.keep_with_next = True
            size, bold = {"Heading 1": (13, True), "Heading 2": (12, True), "Heading 3": (11, True)}.get(name, (11, True))
        else:
            size, bold = 10.5, None
            if re.match(r"^(?:图\d+(?:-\d+)?|表\s*\d+|补充表\d+)\s+", p.text):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                size, bold = 9.5, True
                if p.text.startswith(("表", "补充表")):
                    pf.keep_with_next = True
            elif p.text and not p._p.xpath(".//w:drawing"):
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            font(r, size, bold)
            if name in ('Title', 'Subtitle') or name.startswith('Heading'):
                r.italic = False
        if re.match(r'^\[\d+\]', p.text):
            pf.line_spacing = 1.15
            pf.space_after = Pt(2)
            for r in p.runs:
                font(r, 9.5)
        # 限制图片宽度，按原比例缩小，保留原图字节。
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.keep_with_next = True
        for border in p._p.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    max_width = Mm(166)
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            shape.height = int(shape.height * max_width / shape.width)
            shape.width = max_width
    for t in doc.tables:
        t.autofit = False
        total = sum(c.width or 0 for c in t.columns)
        if total > max_width:
            for row in t.rows:
                for c in row.cells:
                    c.width = int((c.width or max_width / len(row.cells)) * max_width / total)
            for c in t.columns:
                if c.width:
                    c.width = int(c.width * max_width / total)
        pr = t._tbl.tblPr
        borders = pr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            pr.append(borders)
        for tag in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            el = OxmlElement(f"w:{tag}")
            for k, v in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                el.set(qn(f"w:{k}"), v)
            borders.append(el)
        for ri, row in enumerate(t.rows):
            for h in row._tr.xpath("./w:trPr/w:trHeight"):
                h.getparent().remove(h)
            if ri == 0:
                h = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(h)
            # 一行不能跨页拆成两半，表本身允许跨页并重复表头。
            if not row._tr.xpath("./w:trPr/w:cantSplit"):
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for c in row.cells:
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = c._tc.get_or_add_tcPr()
                mar = OxmlElement("w:tcMar")
                for edge in ["top", "left", "bottom", "right"]:
                    e = OxmlElement(f"w:{edge}")
                    e.set(qn("w:w"), "70")
                    e.set(qn("w:type"), "dxa")
                    mar.append(e)
                cp.append(mar)
                if ri == 0:
                    sh = OxmlElement("w:shd")
                    sh.set(qn("w:fill"), "E7E6E6")
                    cp.append(sh)
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = len(t.rows) <= 8 and ri < len(t.rows) - 1
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        font(r, 9.5, ri == 0)
    doc.core_properties.author = author
    doc.core_properties.last_modified_by = author
    doc.core_properties.title = "Wikidata电影知识图谱问答实验报告" if author != "周新永 夏俊杰" else "双人分工与复现实验计划"


def add_latest(doc, person, summary):
    env = anchor(doc, "3 数据集")
    before(doc, env, "2026年10月31日的扩展推理另用H盘项目内.venv（Python 3.11.7）执行；本机GPU为RTX 3070 Ti Laptop 8GiB。自由问法微调与P0服务在Ubuntu 18.04服务器运行，使用项目独立Python 3.11.18、torch 2.6.0+cu118、驱动530.30.02及4张V100-SXM2-32GB中的GPU1。用户允许共享一张卡，分配器限制25%、CPU线程2；不停止其他用户任务。主实验3050训练、后续本机复核与V100推理的环境及计时分别记录。")
    before(doc, env, "环境截图待补1张真实操作系统终端图，须显示解释器路径、Python及依赖版本、当前目录。现有浏览器运行图不能替代这项材料。")
    methods = anchor(doc, "4 方法实现")
    data_end = next((p for p in doc.paragraphs if p.text.startswith("3.2 子图")), methods)
    before(doc, data_end, "3.1 后续数据版本与评价边界", "Heading 2")
    before(doc, data_end, "V1四关系主实验保留5000电影、35485三元组与58994道模板问答不变。V2目录另存5000电影的170516条带类型声明，覆盖12属性；其实体、日期、金额和限定信息由工具执行器按类型读取，不直接充当B0的新训练标签。四关系自由问法v2另有86974题，训练/验证/测试为73202/6868/6904，继承电影划分并按问法family区分；这些问题仍由程序化改写生成。")
    before(doc, data_end, "439题扰动诊断、48题免训练开发验收、20题P0开发验收与8917题V1主测试集分别统计，不能混成同一测试集，也不能直接相减解释为开放域提升。")
    result = anchor(doc, "5 实验结果")
    before(doc, result, "4.4 免训练扩展与P0语义规划", "Heading 2")
    before(doc, result, "先保留A0/B0原始分支，再增加中英文输入适配、范围保护、证据核验和显式工具回退。V2扩展工具D0提供规则查询；P0使用冻结Qwen2.5-3B-Instruct，将问题转为kgqa-plan/1 JSON计划，校验电影候选、属性白名单与预算后，由Explorer.execute确定性执行。当前算子为事实查询、实体交集、导演作品反查与同口径金额比较。")
    before(doc, result, "P0不再把整句问题交给四类关系分类器，而是分开处理片名定位、查询规划、图谱执行和证据回答。模糊片名要求确认；遗漏电影、未执行条件或模型服务故障明确返回澄清、拒答或错误。长列表、日期金额等使用完整模板兜底并记录fallback_used；保留模型原始计划、原始文本、执行结果与证据。权重冻结，工具事实不写回模型。")
    before(doc, result, "仍不支持任意自然语言问题。剧情、主观推荐、开放域评论、任意过滤/计数和未实现的推理链需要新数据或新算子；已有12属性也取决于本地声明覆盖。接入Transformer指令模型扩大了表达入口，不等于已经正确解析所有表达。")
    h52 = anchor(doc, "5.2 基础实验结果")
    before(doc, h52, "P0补充运行记录", "Heading 3")
    before(doc, h52, "2026年10月31日实际下载并启动Qwen/Qwen2.5-3B-Instruct，修订号aa8e72537993ba99e69dfaafa59ed015b17504d1，FP16冻结推理，服务只监听服务器127.0.0.1，通过SSH转发到本机7870端口。GPU1实测进程约2.1—2.5GiB；25%是PyTorch分配器上限，不代表整卡独占。真实问题“Who put together the screenplay for Inception?”执行P58并返回Christopher Nolan；“what else did the person who made it direct? Just three titles.”执行导演反查，记录limit=3及模板回退。这两例用于确认真实HTTP链路，完整开发结果另表列出。")
    before(doc, h52, "P0运行入口为python -m kgqa.semantic_qa；Web新增/api/semantic，并将P0与A0/B0及V2规则工具独立展示。服务、模型清单和逐题输出保存于verification/2026-10-31/semantic_p0/，原数据与旧训练目录保持冻结。")
    h53 = anchor(doc, "5.3 参数修改与重复实验")
    before(doc, h53, "P0开发验收结果", "Heading 3")
    before(doc, h53, P0_LIMIT)
    before(doc, h53, "补充表1 D0与P0开发验收结果")
    rows = []
    for label, key in [("计划匹配", "declared_plan_match"), ("工具结果匹配", "tool_result_match"), ("状态匹配", "status_match"), ("结构任务匹配", "structured_task_match"), ("模板回退", "fallback_rate")]:
        vs = []
        for name in ["D0", "P0"]:
            v = summary[name][key]
            vs.append(f"{v['numerator']}/{v['denominator']}  {v['percent']:.2f}%")
        rows.append([label, *vs])
    rows += [["平均单题延迟 ms", f"{summary['D0']['mean_latency_ms']:.2f}", f"{summary['P0']['mean_latency_ms']:.2f}"], ["P95单题延迟 ms", f"{summary['D0']['p95_latency_ms']:.2f}", f"{summary['P0']['p95_latency_ms']:.2f}"]]
    table_before(doc, h53, ["指标", "D0规则工具", "P0冻结模型规划"], rows, [58, 54, 54])
    before(doc, h53, "延迟为逐题串行墙时，排除初始化、未预热，包含本机到服务器接口开销。计划/结果只在16道有可评价计划的题上计分，状态与结构任务对全部20题计分。原始输出、参考、来源哈希与source_unchanged/inputs_unchanged检查均保留。")
    h54 = anchor(doc, "5.4 基础实验与参数修改实验的结果对比")
    before(doc, h54, "P0本轮只固定一组模型与推理配置，未进行三次学习率训练，也未把开发提示修改冒称课程参数重复实验。本报告的参数实验仍采用上文A0阈值或B0—B3独立学习率运行；P0是另列的功能扩展。")
    h55 = anchor(doc, "5.5 结果对比分析")
    before(doc, h55, "免训练增强与自由问法微调补充对照", "Heading 3")
    before(doc, h55, "48道中英文开发验收中，24道可答、12道不可答、12道需要澄清。原始A0在24道可答题实体F1为100%，但不可答拒答为50%、澄清为16.67%；加保护后分别为100%、83.33%，非可答题的自动非拒答比例由50%降至0。该比例依赖拒答标志及规范文本，不能代替人工事实核查。B0原始/保护/混合分支可答F1为94.79%/91.67%/100%，混合分支2/48题使用工具回退。工具修正不能归因于B0自身能力增长。")
    before(doc, h55, "四关系自由问法v2微调已完成3轮（学习率3e-5、有效batch16、seed42），按验证F1选择第3轮。验证总体F1由89.65%到93.20%，可答F1由88.03%到99.53%，拒答却从95.47%降至70.53%。6904题未见family测试实体F1为69.41%、可答F1为72.48%、拒答58.33%；没有旧B0同测试集对照，不能与原模板测试直接作提升差值。总耗时15632.05秒，进程采样峰值5600MiB；未替换界面默认模型。")
    before(doc, h55, "P0与D0的20题结果则显示：P0状态匹配为70%高于D0的60%，但计划43.75%低于75%，结构任务50%低于60%，平均延迟约5.95秒。这不是模型规模增大便必然提升的结果；提示约束、属性语义、多电影绑定与证据执行需要分别验收。")
    conclusion = anchor(doc, "6 实验总结")
    before(doc, conclusion, "免训练语义扩展的收益与限制", "Heading 3")
    before(doc, conclusion, "当前P0已跑通真实模型规划与图谱执行，但20题开发验收不足以支撑“已理解任意自然语言”。较好的拒答状态与较低的计划匹配可以同时出现：系统可能安全地澄清，也可能选择了错误属性、漏掉条件或把反查简化为单跳。模板回退只保证执行结果按规则完整呈现，不能修正错误计划。应保留计划、实体和证据，以便逐题定位错误。")
    before(doc, conclusion, "后续优先开展免训练改进：加强属性定义与正反例提示，加入候选实体完整性检查，对未支持条件显式拒答；固定提示后，由独立人员编写覆盖更多电影和表达的新题集，分别测计划、工具、状态、文本事实和延迟。只有新方案在独立题上可靠改善，才考虑替换默认分支。进一步训练应保留拒答样本与旧模板回归，避免只优化可答题。")
    before(doc, conclusion, "本机完整代码回归716项通过（39.21秒）。协议模拟、Schema、执行器及Web测试证明工程行为，不能代替真实模型自然语言能力评价。")
    ref = anchor(doc, "7 参考文献")
    role = ZHOU_SCOPE if person == "周新永" else XIA_SCOPE
    before(doc, ref, f"最新报告责任范围为：{role}。两人按模块均衡承担报告解释、复现与答辩，约各占一半；共同核对数据划分、评价口径、CI和结论。该分配依据当前交付模块组织，不追溯原始贡献，也不据此声称每人独立执行过全部服务器任务。")
    before(doc, ref, "人工评分仍为空，真实终端环境截图尚需补入。报告完成日期填写2026年10月30日，10月31日扩展按真实日期单列为补充记录。课程代码包不含完整数据和模型，运行资产按manifest另取。")
    before(doc, ref, "补充记录见verification/2026-10-31/下的freeform_v2/completed/、trainfree/runtime_5000_final_v2/与semantic_p0/dev_v1/；P0模型来源：https://huggingface.co/Qwen/Qwen2.5-3B-Instruct。")


def unify_zhou(doc):
    doc.paragraphs[0].text = "Wikidata电影知识图谱问答实验报告"
    doc.paragraphs[1].text = "人工智能 I 课程实验  问答规划与评价"
    caption = anchor(doc, "表1 实验基本信息")
    before(doc, caption, "1 实验基本信息与个人任务", "Heading 1")
    remove(anchor(doc, "1 实验目的与个人任务"))
    replace(doc, "本实验围绕Wikidata", "本实验围绕Wikidata电影知识图谱问答展开，先在5000电影固定图谱上比较管道与神经生成，再进行阈值、学习率、结构前缀和鲁棒性分析，最后扩展中英文输入、证据保护及冻结指令模型规划。我的报告重点为问答、计划执行与评价；A0基础运行加载配套冻结分类器，本轮未重新训练A0。")
    replace(doc, "我与夏俊杰在同一", f"按当前完整交付物分配，我负责{ZHOU_SCOPE}。夏俊杰负责{XIA_SCOPE}。两人共享数据划分、事实范围及严格评价接口，共同整理结果、CI与报告。分工按最新模块责任组织，工作量约各一半，不追溯历史贡献。")
    replace(doc, "3 数据集与数据划分", "3 数据集与实验设置")
    replace(doc, "4 总体实验设置", "4 方法实现与评价协议")
    replace(doc, "夏俊杰完成TransE的100轮训练", "共同神经对照由数据与模型模块负责，记录包括TransE的100轮训练，以及B0—B3和C0各5轮神经训练。五组使用相同数据与预算：训练batch为8，梯度累积2次，最大有效batch为16；评价batch为64，输入最多512 token，目标最多96 token，随机种子和打乱种子均为42。B0—B3仅改变学习率，C0在B0学习率下移除结构前缀。每轮完整评价8857道验证题，恢复验证选择的检查点后再生成测试答案。")
    replace(doc, "为比较另一类参数变化", "共同神经对照用于比较另一类参数变化。B0—B3从相同预训练模型独立训练，学习率依次为1×10⁻⁴、3×10⁻⁵、3×10⁻⁴、1×10⁻³；C0保持1×10⁻⁴并移除结构前缀。这些结果与A0阈值实验分别解释，不把它们写成A0重新训练。")
    replace(doc, "B0—B3比较不同学习率的实体F1", "B0—B3比较不同学习率的实体F1；C0作为相同学习率下的独立消融项标出。图中分数来自统一实体解析v2，未将名称恢复结果混入原始模型比较。")
    replace(doc, "通过本次实验，我把阈值", "本实验把输入、链接、关系、事实与评分串联起来。阈值0.99以下不改变模板成绩，是因门槛未触及预测置信度；0.998和0.999触发误拒答。扰动诊断说明片名与表达变化仍会导致失败，模板满分不能代表开放域能力。实体解析修订、名称恢复与模型训练分别记录，避免把评分或后处理变化误称模型提升。")
    remove(anchor(doc, "扰动题表明，关系表达"))
    remove(anchor(doc, "中文适配把这些认识"))
    doc.tables[0].add_row().cells[0].text = "报告责任分工"
    doc.tables[0].rows[-1].cells[1].text = ZHOU_SCOPE
    para(doc, "复现时先核对上述数据和模型哈希，再使用新的输出目录。主实验A0预测、严格评分、名称恢复和P0规划各自保留结果；不得覆盖已用于报告的冻结记录。下列命令用于复现说明，不表示本轮又进行了新的选模。")
    para(doc, r".\.venv\Scripts\python.exe benchmarks/run_frozen_thresholds.py --data-dir data/processed/movies_5000_20261029 --model-path models/pipeline_5000_20261029/model.pkl --output-dir output/report_recheck_NEW")
    para(doc, "阈值入口会加载同一冻结模型并分别保存配置、逐题预测与汇总。复核时检查问题ID唯一、8917题完整覆盖、可答8106题与拒答811题的分母，并核对阈值拒答是否发生在关系预测之后。高阈值是诊断实验，不能用已查看的测试数据重新选择正式默认值。")
    para(doc, r".\.venv\Scripts\python.exe scripts/evaluate_semantic_dev.py --catalog-dir data/processed/movies_v2_20261030 --cases benchmarks/semantic_dev_v1.jsonl --output-dir output/semantic_recheck_NEW")
    para(doc, "P0复跑前需检查冻结模型服务与SSH转发；服务不可用应报错，不能切换规则答案冒称P0。20题复跑仅检查同题同协议的可重复行为，仍不是新的独立测试。独立自然问题与人工文本事实评价需另建题集及评分记录。")


def unify_xia(doc):
    # 保留原基本信息表，去掉旧封面与手工导航，将其接到统一标题后。
    table = doc.tables[0]._tbl
    env = anchor(doc, "三 实验环境")
    prefix = []
    for el in list(doc.element.body):
        if el is env._p:
            break
        if el is not table:
            prefix.append(el)
    for el in prefix:
        doc.element.body.remove(el)
    p = before(doc, env, "Wikidata电影知识图谱问答实验报告", "Title")
    table.addprevious(p._p)
    p = before(doc, env, "人工智能 I 课程实验  数据表示与神经生成", "Subtitle")
    table.addprevious(p._p)
    p = before(doc, env, "1 实验基本信息与个人任务", "Heading 1")
    table.addprevious(p._p)
    p = before(doc, env, "表1 实验基本信息")
    table.addprevious(p._p)
    before(doc, env, "本实验在5000电影固定知识图谱上实现A0管道与B0神经生成，对比学习率和结构前缀，随后检查自由问法微调与免训练增强。我的报告以数据、结构表示、生成模型及部署复现为主，同时记录P0冻结模型的真实运行和局限。")
    before(doc, env, f"按当前完整交付物分配，我负责{XIA_SCOPE}。周新永负责{ZHOU_SCOPE}。两人共同核对数据划分、评价、CI、结果和报告，按最新模块责任均衡组织，约各占一半，不追溯历史贡献。共享运行按日志说明，不声称每人独立执行过全部任务。")
    env.text = "2 实验环境"
    replace(doc, "四 数据集与实验设置", "3 数据集与实验设置")
    replace(doc, "4.1 子图构建与数据质量", "3.2 子图构建与数据质量")
    trans = anchor(doc, "4.2 TransE 结构表示")
    before(doc, trans, "4 方法实现与评价协议", "Heading 1")
    trans.text = "4.1 TransE结构表示"
    replace(doc, "4.3 事实文本与结构前缀融合", "4.2 事实文本与结构前缀融合")
    replace(doc, "4.4 训练与评价协议", "4.3 训练与评价协议")
    replace(doc, "五 实验结果", "5 实验结果")
    replace(doc, "六 实验总结", "6 实验总结")
    p = replace(doc, "参考文献与复核资料", "7 参考文献")
    p.style = "Heading 1"
    replace(doc, "5.4 基础与参数修改实验的结果对比", "5.4 基础实验与参数修改实验的结果对比")
    replace(doc, "夏俊杰负责B0—B3学习率对照", "B0—B3为独立学习率对照，C0为结构前缀消融；各组均从同一预训练模型独立训练。数据与模型模块负责配置、曲线和运行资产，问答与评价模块核对实体评分和输出，以下结果共同用于两份报告的不同分析重点。")
    replace(doc, "周新永负责答案实体解码", "答案实体解码与名称恢复的评价接口归入周新永的问答评价模块，模型训练归入我的数据与神经模块。名称恢复作为独立后处理，不重新训练；仅使用实际链接电影的保留事实名称，不读取标准关系或答案进行恢复。原始验证F1仍决定B模型选择，恢复结果不参与选模，也不替换原始成绩。结构消融始终比较原始B0与C0。")
    remove(anchor(doc, "本人的负责范围："))
    replace(doc, "两张实际浏览器模型测试截图", "两张实际浏览器模型测试截图已移入5.1，图号为图5-1、图5-2。人工评价计划尚待两位评价者独立完成，空白评分表表示未评；还需补入1张真实操作系统终端环境截图，显示实际Python、依赖版本和项目运行目录。")
    # 实际浏览器运行图从5.4移入5.1，保留图像、关系及分辨率。
    h52 = anchor(doc, "5.2 基础实验结果")
    for start in ["图5-7 5000部电影双系统", "图5-8 特殊字符"]:
        cap = anchor(doc, start)
        image_el = cap._p.getprevious()
        h52._p.addprevious(image_el)
        h52._p.addprevious(cap._p)
    before(doc, h52, "图5-1和图5-2为实际浏览器模型测试，分别记录双系统问答及特殊字符名称恢复。它们说明加载模型、输入问题、生成答案与展示证据的流程；不以单例截图替代8917题批量评价。训练过程的完整history、config和逐题记录保留在verification/2026-10-30/。")
    mapping = {"1": "3", "2": "4", "3": "5", "4": "6", "5": "7", "6": "8", "7": "1", "8": "2"}
    for p in doc.paragraphs:
        if re.match(r"^图5-[1-8] ", p.text):
            p.text = re.sub(r"^图5-([1-8])", lambda m: "图5-" + mapping[m[1]], p.text)
            if p.text.startswith('图5-5 六种系统'):
                p.text = '图5-5 六种基础系统及名称恢复的完整测试比较'
    # 修复原稿表号重复，主实验表按出现顺序编号。
    count = 0
    for p in doc.paragraphs:
        if re.match(r"^表\s*\d+", p.text):
            count += 1
            p.text = re.sub(r"^表\s*\d+\s*", f"表{count} ", p.text)
    for row in doc.tables[0].rows:
        if row.cells[0].text == "主要分工":
            row.cells[1].text = XIA_SCOPE
        if row.cells[0].text == "整理日期":
            row.cells[1].text = "2026年10月31日"


def build_plan(old, summary):
    doc = Document(old)
    # 分工计划过时，按当前成果重写；保留原件备份。
    for el in list(doc.element.body):
        if el.tag != qn("w:sectPr"):
            doc.element.body.remove(el)
    para(doc, "双人分工与复现实验计划", "Title")
    para(doc, "人工智能 I 课程实验  2026年10月31日", "Subtitle")
    para(doc, "本计划按最新模块与可复核交付物分配两人的报告、复现和答辩责任，工作量约各一半，不追溯原始贡献。项目已完成5000电影主实验与若干扩展；P0真实链路已运行，但开发验收仍暴露语义规划不足。两份个人报告使用统一结构，各自深入不同实验重点，共享事实及指标口径。")
    para(doc, "1 成员与均衡分工", "Heading 1")
    end = para(doc, "各成员对负责模块的原理、配置、逐题结果和局限承担解释与复现责任。共享日志不据此改写为个人独立执行历史，既有冻结数据、旧模型和同伴任务保持不变。")
    table_before(doc, end, ["成员", "责任模块", "交付与报告重点"], [
        ["周新永\nM202674502\n硕2609班", "A0与关系阈值；电影链接、候选确认；P0计划校验与执行；评价与错误诊断；证据和澄清交互", "问答流程和阈值三次修改；严格实体评价；439题鲁棒性；48题保护/回退；20题D0/P0开发验收"],
        ["夏俊杰\nM202674534\n硕2610班", "V1/V2数据来源；TransE；B0—B3/C0；自由问法微调；冻结Qwen服务；生成、脚本和环境", "数据/模型参数与学习率三次修改；结构前缀消融；训练曲线；服务器资源与部署；生成和模板回退"],
        ["共同", "划分与种子；实验协议；CI与复现；结果核查；报告与答辩", "统一系统命名、分母、时间范围和结论；核对课程提交材料；独立人工评分如实补做"],
    ], [28, 69, 69])
    para(doc, "工作量按模块复杂度和交付责任平衡，接口整合及最终复核共同完成。两份报告的完成日期均为2026年10月30日，10月31日扩展按实际日期列为补充记录。")
    para(doc, "2 当前实验完成情况", "Heading 1")
    end = para(doc, "数据版本分开：V1为四关系5000电影主实验，V2为12属性工具目录；V2声明不自动变为已训练能力。20题与48题均是程序编写开发验收，不能声称独立自然问答泛化已通过。")
    table_before(doc, end, ["模块", "当前状态", "证据位置"], [
        ["V1数据", "5000电影；7864实体；35485三元组；58994问答；电影3500/750/750，问答41220/8857/8917；seed42", "datasets/movies-5000-20261029/manifest.json"],
        ["A0主实验", "8917题实体F1为100%；0.70/0.90/0.99三次修改无变化；诊断高阈值0.998/0.999触发误拒答", "verification/2026-10-30/A0_thresholds/；reports/补充阈值运行记录/"],
        ["TransE和B/C", "TransE100轮；B0—B3、C0各5轮；验证同分规则选B0，测试F1 99.53%；B0/C0本轮F1相同", "verification/2026-10-30/"],
        ["鲁棒性与恢复", "439题；改写/错拼显著下降；B0独立名称恢复后模板测试F1为100%，扰动无改善", "decoder_revision_v2/；paired_robustness.json"],
        ["V2工具", "5000电影170516条声明；12属性；类型、日期金额及证据；候选确认和支持范围限制", "data/processed/movies_v2_20261030/manifest.json"],
        ["自由问法微调", "3轮已完成；验证F1 93.20%，拒答下降；6904题测试F1 69.41%；未切换默认模型", "verification/2026-10-31/freeform_v2/completed/"],
        ["免训练保护/回退", "48题开发验收；保护后拒答100%、澄清83.33%；混合B0有2题工具回退，单列来源", "verification/2026-10-31/trainfree/runtime_5000_final_v2/"],
        ["P0冻结规划", "真实Qwen2.5-3B FP16服务；20题开发验收完成，权重未更新；计划匹配仍不足", "verification/2026-10-31/semantic_p0/"],
        ["工程回归", "本机716 passed，39.21秒；协议模拟、执行器、Web和回归分别验证", "实际完整测试输出及tests/；不等于模型准确率"],
    ], [25, 88, 53])
    para(doc, "3 P0补充结果与结论", "Heading 1")
    para(doc, P0_LIMIT)
    end = para(doc, "P0状态处理在这批开发题上较好，但计划与结构任务匹配低于D0，且交互延迟较大。不能写成已完整理解自然语言、替代B0或提升原主实验分数。后续先固定提示和校验，再收集独立新题；生成句事实与流畅度仍待人工复核。")
    table_before(doc, end, ["指标", "D0规则工具", "P0冻结模型"], [
        ["计划匹配", "12/16  75.00%", "7/16  43.75%"], ["工具结果匹配", "8/16  50.00%", "7/16  43.75%"],
        ["状态匹配", "12/20  60.00%", "14/20  70.00%"], ["结构任务匹配", "12/20  60.00%", "10/20  50.00%"],
        ["平均/P95 ms", "15.80 / 26.30", "5954.60 / 16901.10"], ["模板回退", "0/20", "11/20"],
    ], [58, 54, 54])
    para(doc, "4 复现顺序与验收", "Heading 1")
    para(doc, "第一步核对Git版本、数据manifest、SHA256SUMS、模型和嵌入映射。课程代码包不含完整数据、权重及缓存；完整5000资产按数据manifest和独立运行包取得。使用H盘项目内.venv，不覆盖冻结目录。A0/B0原始、名称恢复、保护/回退、D0和P0各保存独立预测和运行名。")
    para(doc, "第二步执行工程检查，由两人共同复核失败与跳过项。以下命令从H:\\AI1experiment运行：")
    para(doc, "静态检查需ruff 0.12.12（见requirements-ci.txt）。本轮使用H盘既有CI环境运行；若主.venv未安装，先执行python -m pip install ruff==0.12.12再运行下列ruff命令，避免误将依赖缺失当成检查通过。")
    for command in [r".\.venv\Scripts\python.exe -m pytest -q", r".\.venv\Scripts\python.exe -m ruff check --select E9,F63,F7,F82 kgqa tests", r".\.venv\Scripts\python.exe -m compileall -q kgqa tests", r"git diff --check"]:
        para(doc, command)
    para(doc, "第三步由周新永复核A0阈值和严格评分，夏俊杰复核B/C配置、验证选模和数据模型配套，之后交叉核对逐题ID、分母、错误类型与计时范围。A0全流程单题延迟、神经batch生成摊销和P0串行HTTP墙时不能直接计算加速倍数。已查看测试分数后的高阈值诊断不作为独立寻优。")
    para(doc, "第四步启动冻结模型服务并做P0开发复跑。Qwen修订号aa8e72537993ba99e69dfaafa59ed015b17504d1；远端只监听127.0.0.1，按docs/服务器与训练运行说明.md检查GPU并保持25%分配器限额、CPU线程2。现有服务和训练目录先检查，避免重复启动或覆盖。")
    para(doc, r"ssh -N -L 7870:127.0.0.1:7870 -p 50012 junjiex@222.20.94.68")
    para(doc, r".\.venv\Scripts\python.exe -m kgqa.semantic_qa --help")
    para(doc, r".\.venv\Scripts\python.exe scripts/evaluate_semantic_dev.py --help")
    para(doc, "按--help指定V2目录、冻结题集与新的空输出目录复跑，原开发结果不覆盖。命令帮助只说明入口可用，不能记成模型推理成功。真实服务清单与20题预测已另存；重新测量时记录网络和共享GPU负载。")
    para(doc, "第五步验收界面：分别检查A0/B0、V2工具、P0模式；查看电影候选确认、完整计划、原始模型输出、证据、拒答和fallback标识。模型不可用应显示服务错误，不能悄悄以规则答案冒充P0。")
    para(doc, "5 报告与课程要求对应", "Heading 1")
    end = para(doc, "两份个人报告均按1基本信息、2环境、3数据、4实现、5.1—5.5结果与分析、6总结、7参考资料组织，正文与图表格式统一，原有真实图像和公式保留。不同人的核心参数实验和讨论重点保持区别。")
    table_before(doc, end, ["课程要求", "周新永报告", "夏俊杰报告"], [
        ["基础设置和结果", "冻结A0、8917题、实体/文本指标及管道耗时", "TransE及B0，完整训练/验证/测试和生成指标"],
        ["至少三次参数修改", "0.70、0.90、0.99；另列诊断高阈值", "B1 3e-5、B2 3e-4、B3 1e-3，相对B0 1e-4"],
        ["对比图与分析", "阈值图及神经对照；误拒答/链接/评测", "学习率损失与指标、消融、鲁棒性和TransE曲线"],
        ["至少两张运行图", "原有英/中文实际浏览器测试", "原有双系统及名称恢复实际浏览器测试"],
        ["终端环境图", "仍待补入真实终端图", "仍待补入真实终端图"],
        ["人工评价", "自设计划尚未评分，不虚构分数", "自设计划尚未评分，不虚构分数"],
    ], [38, 64, 64])
    para(doc, "6 待完成事项与改进顺序", "Heading 1")
    para(doc, "提交前仍需补入真实终端环境截图；两份已有实际浏览器测试图保留上下文并放入运行记录。人工评分是本项目原定评价计划，尚未执行；教师所列BLEU、ROUGE及人工评估为示例，不虚构评分。")
    para(doc, "人工计划采用相同100题、A0与验证选定B系统各一份答案，两名评价者独立填写准确性0—2、完整性0—2、流畅度1—5和无依据断言0/1，原始表每人200条。评分空白继续表示未评，原始分歧与讨论另存，不用自动重叠指标或LLM评分代填。")
    return doc


def fill_identity(doc, person):
    values = {'学号': 'M202674502' if person == '周新永' else 'M202674534',
              '班级': '硕2609班' if person == '周新永' else '硕2610班',
              '实验完成日期': '2026年10月30日', '课程实验完成日期': '2026年10月30日'}
    for row in doc.tables[0].rows:
        if row.cells[0].text in values:
            row.cells[1].text = values[row.cells[0].text]
    if person == '夏俊杰':
        row = doc.tables[0].add_row()
        row.cells[0].text, row.cells[1].text = '补充运行日期', '2026年10月31日'


def package_stats(path):
    import zipfile
    with zipfile.ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--reports-dir', type=Path, default=ROOT / 'reports')
    ap.add_argument('--summary', type=Path, default=ROOT / 'verification/2026-10-31/semantic_p0/dev_v1/summary.json')
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()
    summary = json.loads(args.summary.read_text(encoding='utf-8'))
    for s in ['D0', 'P0']:
        assert summary[s]['count'] == 20
    paths = [args.reports_dir / f for f in FILES]
    for p in paths:
        if not p.is_file():
            raise FileNotFoundError(p)
        if any('D0与P0开发验收结果' in para.text for para in Document(p).paragraphs):
            raise ValueError('已更新的文件不重复插入，请使用备份原件或单独修订')
    backup = ROOT / '.work/semantic-report-backups' / datetime.now().strftime('%Y%m%d_%H%M%S')
    print(json.dumps({'inputs': [str(p) for p in paths], 'backup': str(backup), 'dry_run': args.dry_run, 'p0_plan': summary['P0']['declared_plan_match']}, ensure_ascii=False, indent=2))
    if args.dry_run:
        return
    backup.mkdir(parents=True, exist_ok=False)
    for p in paths:
        shutil.copy2(p, backup / p.name)
    records = []
    for p, name in zip(paths[:2], ['周新永', '夏俊杰']):
        orig_media = package_stats(p)
        d = Document(p)
        math_count = len(d.element.xpath('.//m:oMath'))
        (unify_zhou if name == '周新永' else unify_xia)(d)
        add_latest(d, name, summary)
        fill_identity(d, name)
        format_doc(d, name)
        if name == '周新永':
            anchor(d, '附录 实验记录索引').paragraph_format.page_break_before = True
        tmp = p.with_name(p.stem + '.updating.docx')
        d.save(tmp)
        if package_stats(tmp) != orig_media:
            raise ValueError(f'原图像发生变化 {p}')
        if len(Document(tmp).element.xpath('.//m:oMath')) != math_count:
            raise ValueError(f'原生公式数量发生变化 {p}')
        tmp.replace(p)
        records.append({'path': str(p), 'sha256': sha(p), 'original_media_preserved': True, 'native_math_count': math_count, 'paragraph_count': len(d.paragraphs), 'table_count': len(d.tables), 'inline_shapes': len(d.inline_shapes)})
    d = build_plan(paths[2], summary)
    format_doc(d, '周新永 夏俊杰')
    d.save(paths[2])
    records.append({'path': str(paths[2]), 'sha256': sha(paths[2])})
    (backup / 'operation.json').write_text(json.dumps({'backup': str(backup), 'outputs': records}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(records, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
